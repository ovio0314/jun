import io
from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from voucher.export import safe_cell, to_excel_bytes


def test_safe_cell():
    assert safe_cell("=SUM(A1)") == "'=SUM(A1)"
    assert safe_cell("+1") == "'+1"
    assert safe_cell("@cmd") == "'@cmd"
    assert safe_cell("원재료") == "원재료"
    assert safe_cell(Decimal("1000")) == 1000
    assert safe_cell(Decimal("0.5")) == "0.5"


def test_excel_has_no_formula():
    data = to_excel_bytes({"시트": pd.DataFrame({"적요": ["=HYPERLINK(\"x\")", "정상"], "금액": [Decimal(5), Decimal(7)]})})
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A2"].data_type != "f" and ws["A2"].value.startswith("'=")
    assert ws["B2"].value == 5
