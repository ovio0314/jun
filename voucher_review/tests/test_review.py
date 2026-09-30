import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from make_synthetic import build_rows, write_xlsx
from voucher.checks import run_checks, summary_counts, voucher_summary
from voucher.evidence import ManualEvidenceStore
from voucher.export import build_result_sheets, check_export_counts, to_excel_bytes
from voucher.parser import find_header_candidates, load_vouchers
from voucher.review import (
    content_hashes,
    delete_save,
    init_reviews,
    load_session,
    reviews_frame,
    save_session,
    update_review,
)
from voucher.settings import default_settings
from ui.review_page import apply_filters


def _load(path):
    return load_vouchers(path, find_header_candidates(path)[0])


def test_reupload_changed_voucher_loses_completed(tmp_path):
    rows = build_rows(n_vouchers=3)
    first = _load(write_xlsx(tmp_path / "a.xlsx", rows=rows))
    reviews = init_reviews(content_hashes(first.df))
    keys = list(reviews)
    for k in keys:
        update_review(reviews, k, status="검토완료", memo="확인함", reviewer="검토자1")

    rows2 = [dict(r) for r in rows]
    target = rows2[0]["기표번호"]
    rows2[0]["적요"] = "원재료 매입 (수정)"  # 첫 전표 분개 변경
    second = _load(write_xlsx(tmp_path / "b.xlsx", rows=list(reversed(rows2))))  # 행 순서 변경은 변경 아님
    carried = init_reviews(content_hashes(second.df), reviews)
    changed = {k: r for k, r in carried.items() if r.변경감지}
    assert list(changed) == [f"본사|회계|{target}"]
    assert changed[f"본사|회계|{target}"].검토상태 == "미검토"
    assert changed[f"본사|회계|{target}"].이전검토상태 == "검토완료"
    assert changed[f"본사|회계|{target}"].메모 == "확인함"
    assert all(r.검토상태 == "검토완료" for k, r in carried.items() if k not in changed)


def test_update_review_validates():
    reviews = init_reviews({"k": "h"})
    with pytest.raises(ValueError):
        update_review(reviews, "k", status="승인")


def test_export_counts_and_formula_injection(tmp_path):
    rows = build_rows(n_vouchers=5)
    rows[0]["적요"] = '=HYPERLINK("http://x","클릭")'
    rows[1]["적요"] = "+SUM(1,2)"
    res = _load(write_xlsx(tmp_path / "e.xlsx", rows=rows))
    results = run_checks(res.df, default_settings())
    summ = voucher_summary(res.df, results)
    reviews = init_reviews(content_hashes(res.df))
    update_review(reviews, summ["전표키"].iloc[0], memo="@메모 수식 아님")
    crit = pd.DataFrame([{"항목": "x", "값": "=1+1"}])
    sheets = build_result_sheets(summ, results, res.raw, res.df, reviews_frame(reviews), crit)
    assert check_export_counts(sheets, len(res.df), len(summ)) == []
    wb = load_workbook(io.BytesIO(to_excel_bytes(sheets)))
    assert set(wb.sheetnames) == {"전표요약", "검사상세", "원본분개", "적용기준"}
    assert wb["원본분개"].max_row - 1 == len(res.df)
    assert wb["전표요약"].max_row - 1 == len(summ)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
    memo_col = [c.value for c in wb["전표요약"][1]].index("메모")
    assert wb["전표요약"].cell(2, memo_col + 1).value == "'@메모 수식 아님"


def test_export_count_mismatch_detected(loaded):
    results = run_checks(loaded.df, default_settings())
    summ = voucher_summary(loaded.df, results)
    reviews = reviews_frame(init_reviews(content_hashes(loaded.df)))
    sheets = build_result_sheets(summ, results, loaded.raw, loaded.df, reviews, pd.DataFrame())
    sheets["원본분개"] = sheets["원본분개"].iloc[:-1]
    assert check_export_counts(sheets, len(loaded.df), len(summ))


def test_filters_do_not_double_count(loaded):
    results = run_checks(loaded.df, default_settings())
    summ = voucher_summary(loaded.df, results)
    reviews = reviews_frame(init_reviews(content_hashes(loaded.df)))
    all_view = apply_filters(summ, reviews, {})
    assert len(all_view) == len(summ) and all_view["전표키"].is_unique
    dept = summ["기표부서"].iloc[0]
    part = apply_filters(summ, reviews, {"기표부서": [dept]})
    rest = apply_filters(summ, reviews, {"기표부서": sorted(set(summ["기표부서"]) - {dept})})
    assert len(part) + len(rest) == len(summ)
    by_result = sum(len(apply_filters(summ, reviews, {"검사결과": [s]})) for s in summary_counts(summ) if s != "전체")
    assert by_result == len(summ)


def test_save_reopen_roundtrip_and_delete(tmp_path, loaded):
    reviews = init_reviews(content_hashes(loaded.df))
    k = next(iter(reviews))
    update_review(reviews, k, status="검토중", memo="메모")
    folder = tmp_path / "saves"
    p = save_session(folder / "s.json", source_name="합성.xlsx", source_sha256="abc", sheet="Sheet1", header_row=2,
                     raw=loaded.raw, df=loaded.df, reviews=reviews, evidence_json=ManualEvidenceStore().dumps(),
                     settings=default_settings(), preamble=["전표조건검색"])
    data = load_session(p)
    assert len(data["df"]) == len(loaded.df)
    assert content_hashes(data["df"]) == content_hashes(loaded.df)
    assert data["review_states"][k].검토상태 == "검토중"
    with pytest.raises(ValueError):
        delete_save(tmp_path / "other.json", folder)
    delete_save(p, folder)
    assert not p.exists()


def test_blank_date_voucher_visible_without_date_filter(tmp_path):
    rows = build_rows(n_vouchers=3)
    for r in rows[:3]:
        r["회계일"] = None
    res = _load(write_xlsx(tmp_path / "nd.xlsx", rows=rows))
    results = run_checks(res.df, default_settings())
    summ = voucher_summary(res.df, results)
    reviews = reviews_frame(init_reviews(content_hashes(res.df)))
    assert len(apply_filters(summ, reviews, {})) == 3
    assert (summ.set_index("기표번호").loc[rows[0]["기표번호"], "데이터 점검"]) == "오류"  # D01 회계일 누락
