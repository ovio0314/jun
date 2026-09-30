"""전표 엑셀 파서.

- 여러 시트에서 헤더 후보를 찾는다.
- 원본 값(raw)과 정규화 값(norm)을 분리해 추적한다.
- 금액은 Decimal 로 변환하고, 해석 불가능한 값은 0 으로 대체하지 않고 오류로 남긴다.
- 원본 파일은 읽기만 한다.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
from openpyxl import load_workbook

from .columns import (
    AMOUNT_COLUMNS,
    DATE_COLUMNS,
    EXPECTED_COLUMNS,
    FLAG_COLUMNS,
    HEADER_SCAN_ROWS,
    MIN_HEADER_MATCH,
    REQUIRED_COLUMNS,
)


class VoucherFileError(Exception):
    """사용자에게 한글로 안내할 파일 오류."""


# 공백류: 일반 공백, 탭, 줄바꿈, NBSP, 전각 공백, 제로폭 공백
_WS_RE = re.compile(r"[\s 　​﻿]+")


def clean_text(value: Any) -> str:
    """앞뒤 공백 제거 + 내부 연속 공백을 하나로. 원본은 따로 보존한다."""
    if value is None:
        return ""
    return _WS_RE.sub(" ", raw_text(value)).strip()


def raw_text(value: Any) -> str:
    """셀 값을 원본에 가깝게 문자열로 보존한다 (숫자로 저장된 코드 포함)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


# ---------------------------------------------------------------- 금액 ----
_NEGATIVE_MARKS = ("△", "▲", "-")


def parse_amount(value: Any) -> tuple[Decimal | None, str]:
    """금액 해석.

    반환: (Decimal 또는 None, 상태)
      상태: "정상" | "빈값" | "오류"
    - 빈값: None / 빈 문자열 / 공백만 → None, "빈값"  (합계 시 0 으로 취급, 상태는 별도 보존)
    - 쉼표, 공백, '원' 접미어는 제거한다.
    - 음수: '-', '△', '▲' 접두 또는 괄호 표기 '(1,000)' 허용. 음수 자체를 오류로 보지 않는다.
    - 그 외 해석 불가능한 값은 None, "오류" (0 으로 대체하지 않음).
    """
    if value is None:
        return None, "빈값"
    if isinstance(value, bool):
        return None, "오류"
    if isinstance(value, int):
        return Decimal(value), "정상"
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None, "오류"
        return Decimal(repr(value)), "정상"
    if isinstance(value, Decimal):
        return value, "정상"

    text = _WS_RE.sub("", str(value))
    if text == "":
        return None, "빈값"
    text = text.replace(",", "")
    if text.endswith("원"):
        text = text[:-1]
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    for mark in _NEGATIVE_MARKS:
        if text.startswith(mark):
            negative, text = (not negative), text[len(mark):]
            break
    if text.startswith("+"):
        text = text[1:]
    if not re.fullmatch(r"\d+(\.\d+)?", text):
        return None, "오류"
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None, "오류"
    return (-amount if negative else amount), "정상"


# ---------------------------------------------------------------- 플래그 ----
def classify_flag(value: Any) -> str:
    """문자열 TRUE/FALSE, 실제 불리언, 빈값을 구분한다."""
    if value is None:
        return "빈값"
    if isinstance(value, bool):
        return "불리언 TRUE" if value else "불리언 FALSE"
    text = str(value).strip()
    if text == "":
        return "빈값"
    if text.upper() == "TRUE":
        return "문자열 TRUE"
    if text.upper() == "FALSE":
        return "문자열 FALSE"
    return f"기타값({text})"


# ---------------------------------------------------------------- 날짜 ----
def parse_date(value: Any) -> dt.date | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = clean_text(value)
    if not text:
        return None
    digits = re.sub(r"[^\d]", "", text)
    if len(digits) == 8:
        try:
            return dt.date(int(digits[:4]), int(digits[4:6]), int(digits[6:]))
        except ValueError:
            return None
    return None


# ---------------------------------------------------------------- 헤더 탐색 ----
@dataclass
class HeaderCandidate:
    sheet: str
    header_row: int  # 1-based 엑셀 행 번호
    matched: int
    title: str  # 헤더 위 첫 텍스트(예: 전표조건검색)
    headers: list[str]
    preamble: list[str] = field(default_factory=list)  # 헤더 위 모든 텍스트(조회 조건·필터 범위)

    @property
    def label(self) -> str:
        return f"{self.sheet} / {self.header_row}행 (일치 열 {self.matched}개)"


def _open_workbook(source: str | Path | bytes | BinaryIO):
    try:
        if isinstance(source, (bytes, bytearray)):
            source = io.BytesIO(source)
        return load_workbook(source, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise VoucherFileError(
            "엑셀 파일을 읽을 수 없습니다. .xlsx 형식인지, 암호가 걸려 있지 않은지 확인해 주세요. "
            f"(상세: {type(exc).__name__})"
        ) from exc
    except Exception as exc:  # openpyxl 내부 오류 등
        raise VoucherFileError(f"엑셀 파일을 여는 중 오류가 발생했습니다: {type(exc).__name__}") from exc


def find_header_candidates(source) -> list[HeaderCandidate]:
    """모든 시트의 상단에서 예상 열 이름과 가장 많이 일치하는 행을 찾는다."""
    wb = _open_workbook(source)
    expected = set(EXPECTED_COLUMNS)
    candidates: list[HeaderCandidate] = []
    try:
        for ws in wb.worksheets:
            best: HeaderCandidate | None = None
            title = ""
            preamble: list[str] = []
            for idx, row in enumerate(ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
                headers = [clean_text(v) for v in row]
                matched = len(expected.intersection(h for h in headers if h))
                if matched >= MIN_HEADER_MATCH and (best is None or matched > best.matched):
                    best = HeaderCandidate(ws.title, idx, matched, title, headers, list(preamble))
                if best is None and matched < MIN_HEADER_MATCH:
                    text = " ".join(h for h in headers if h)
                    if text:
                        preamble.append(text)
                        title = title or next(h for h in headers if h)
            if best:
                candidates.append(best)
    finally:
        wb.close()
    candidates.sort(key=lambda c: -c.matched)
    return candidates


# ---------------------------------------------------------------- 로딩 ----
@dataclass
class LoadResult:
    raw: pd.DataFrame  # 원본 문자열 (공백 정리 전)
    df: pd.DataFrame  # 정규화 값
    mapping: dict[str, str]  # 표준 열 → 원본 헤더
    sheet: str
    header_row: int
    title: str
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    preamble: list[str] = field(default_factory=list)  # 헤더 위 조회 조건 텍스트


def auto_mapping(headers: list[str]) -> dict[str, str]:
    """정확히 일치하는 헤더를 표준 열에 자동 매핑 (중복 헤더는 첫 번째)."""
    mapping: dict[str, str] = {}
    for h in headers:
        if h in EXPECTED_COLUMNS and h not in mapping:
            mapping[h] = h
    return mapping


def duplicate_headers(headers: list[str]) -> list[str]:
    seen, dups = set(), []
    for h in headers:
        if not h:
            continue
        if h in seen and h not in dups:
            dups.append(h)
        seen.add(h)
    return dups


def load_vouchers(source, candidate: HeaderCandidate, mapping: dict[str, str] | None = None) -> LoadResult:
    """선택한 시트/헤더 행 기준으로 데이터를 읽는다."""
    headers = candidate.headers
    mapping = dict(mapping or auto_mapping(headers))
    errors: list[str] = []
    warnings: list[str] = []

    dups = duplicate_headers(headers)
    for d in dups:
        if d in mapping.values() and d in REQUIRED_COLUMNS:
            errors.append(f"필수 열 '{d}' 헤더가 중복되어 있습니다. 원본에서 중복 열을 확인해 주세요.")
        else:
            warnings.append(f"헤더 '{d}' 가 중복되어 있어 첫 번째 열만 사용합니다.")

    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    if missing:
        errors.append("필수 열이 없습니다: " + ", ".join(missing) + ". 열 매핑을 확인해 주세요.")
    optional_missing = [c for c in EXPECTED_COLUMNS if c not in mapping and c not in REQUIRED_COLUMNS]
    if optional_missing:
        warnings.append("선택 열 누락(해당 정보는 표시되지 않음): " + ", ".join(optional_missing))

    col_index = {std: headers.index(src) for std, src in mapping.items() if src in headers}

    wb = _open_workbook(source)
    try:
        ws = wb[candidate.sheet]
        raw_rows, norm_rows = [], []
        for excel_row, row in enumerate(
            ws.iter_rows(min_row=candidate.header_row + 1, values_only=True),
            start=candidate.header_row + 1,
        ):
            if row is None or all(v is None or clean_text(v) == "" for v in row):
                continue
            raw_rec: dict[str, Any] = {"_엑셀행": excel_row}
            norm_rec: dict[str, Any] = {"_엑셀행": excel_row}
            for std in EXPECTED_COLUMNS:
                idx = col_index.get(std)
                value = row[idx] if idx is not None and idx < len(row) else None
                raw_rec[std] = raw_text(value)
                norm_rec.update(_normalize_cell(std, value))
            raw_rows.append(raw_rec)
            norm_rows.append(norm_rec)
    finally:
        wb.close()

    raw = pd.DataFrame(raw_rows, dtype=object)
    df = pd.DataFrame(norm_rows, dtype=object)
    if df.empty:
        errors.append("헤더 아래에 데이터 행이 없습니다.")
        raw = pd.DataFrame(columns=["_엑셀행", *EXPECTED_COLUMNS], dtype=object)
        df = pd.DataFrame(columns=["_엑셀행", "전표키", *EXPECTED_COLUMNS], dtype=object)
    else:
        df.insert(1, "전표키", [voucher_key(r) for r in df.to_dict("records")])
        blank_no = int((df["기표번호"] == "").sum())
        if blank_no:
            warnings.append(f"기표번호가 비어 있는 행 {blank_no}건 — 다른 행과 묶지 않고 행별로 따로 표시합니다.")
        bad = df[[f"{c}_상태" for c in ("차변금액", "대변금액")]].eq("오류").any(axis=1)
        if bad.any():
            warnings.append(
                f"원화 금액을 해석할 수 없는 행 {int(bad.sum())}건 — 0 으로 대체하지 않고 합계에서 제외했습니다."
            )
    return LoadResult(raw, df, mapping, candidate.sheet, candidate.header_row, candidate.title, errors, warnings,
                      list(candidate.preamble))


def voucher_key(rec: dict) -> str:
    """전표 키 = 회계단위|전표관리단위|기표번호. 기표번호가 비면 행마다 별도 키(묶지 않음)."""
    no = rec.get("기표번호") or f"(기표번호없음-{rec.get('_엑셀행')}행)"
    return f"{rec.get('회계단위', '')}|{rec.get('전표관리단위', '')}|{no}"


def _normalize_cell(std: str, value: Any) -> dict[str, Any]:
    if std in AMOUNT_COLUMNS:
        amount, status = parse_amount(value)
        return {std: amount, f"{std}_상태": status}
    if std in DATE_COLUMNS:
        return {std: parse_date(value)}
    if std in FLAG_COLUMNS:
        return {std: clean_text(value), f"{std}_구분": classify_flag(value)}
    return {std: clean_text(value)}


def load_default(path: str | Path) -> tuple[list[HeaderCandidate], bytes]:
    data = Path(path).read_bytes()
    return find_header_candidates(data), data
