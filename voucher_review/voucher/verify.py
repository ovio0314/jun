"""실제 파일 건수 대조 도구 (집계값만 출력, 전표 내용은 출력하지 않음).

실행: python -m voucher.verify "input/9월 전표 내역.xlsx"
- openpyxl 로 원본 데이터 행 수를 파서와 별도로 직접 센다.
- 파서·검사·다운로드 결과의 행/전표 건수를 대조하고 차이를 보고한다.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from .accounts import file_overview
from .checks import SUMMARY_ORDER, is_vat_row, run_checks, summary_counts, voucher_summary
from .export import build_result_sheets, check_export_counts
from .parser import find_header_candidates, load_vouchers
from .review import content_hashes, init_reviews, reviews_frame
from .settings import default_settings


def count_source_rows(path, sheet: str, header_row: int) -> tuple[int, int]:
    """(헤더 아래 비어 있지 않은 행 수, 헤더 아래 전체 행 수)"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        nonblank = total = 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            total += 1
            if any(v is not None and str(v).strip() != "" for v in row):
                nonblank += 1
        return nonblank, total
    finally:
        wb.close()


def verify(path, settings: dict | None = None) -> dict:
    settings = settings or default_settings()
    cands = find_header_candidates(path)
    if not cands:
        raise SystemExit("헤더를 찾지 못했습니다.")
    cand = cands[0]
    res = load_vouchers(path, cand)
    report: dict = {"시트": cand.sheet, "헤더행": cand.header_row, "헤더위내용": cand.preamble,
                    "후보수": len(cands), "오류": res.errors, "경고": res.warnings}
    nonblank, total = count_source_rows(path, cand.sheet, cand.header_row)
    report["원본_비어있지않은행"] = nonblank
    report["원본_헤더아래전체행"] = total
    report["파서_행수"] = len(res.df)
    if res.errors:
        return report
    ov = file_overview(res.df)
    results = run_checks(res.df, settings)
    summ = voucher_summary(res.df, results)
    sheets = build_result_sheets(summ, results, res.raw, res.df, reviews_frame(init_reviews(content_hashes(res.df))),
                                 results.head(0))
    vat = [r for r in res.df.to_dict("records") if is_vat_row(r, settings["vat"])]
    report.update({
        "전표수(회계단위+전표관리단위+기표번호)": ov["전표수"], "기표번호수": ov["기표번호수"],
        "회계일": f"{ov['회계일_시작']} ~ {ov['회계일_종료']}",
        "차대변불일치또는추출누락확인": ov["차대변불일치전표수"],
        "부가세행(계정명/코드 기준)": len(vat), "부가세전표": len({r['전표키'] for r in vat}),
        "적요빈칸행": ov["적요빈칸행수"], "파일첨부여부구분": ov["파일첨부여부구분"],
        "요약건수": {k: v for k, v in summary_counts(summ).items()},
        "다운로드_원본분개행": len(sheets["원본분개"]), "다운로드_전표요약건": len(sheets["전표요약"]),
        "다운로드_대조문제": check_export_counts(sheets, len(res.df), ov["전표수"]),
    })
    diffs = []
    if nonblank != len(res.df):
        diffs.append(f"원본 비어있지 않은 행 {nonblank} ≠ 파서 {len(res.df)}")
    if sum(report["요약건수"][s] for s in SUMMARY_ORDER) != ov["전표수"]:
        diffs.append("요약 4상태 합계 ≠ 전표 수")
    report["차이"] = diffs + report["다운로드_대조문제"]
    return report


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print('사용법: python -m voucher.verify "input/9월 전표 내역.xlsx"')
        return 2
    path = Path(argv[1])
    if not path.exists():
        print(f"파일이 없습니다: {path}")
        return 2
    report = verify(path)
    for k, v in report.items():
        print(f"{k}: {v}")
    print("결과:", "차이 없음" if not report.get("차이") and not report["오류"] else "확인 필요")
    return 0 if not report.get("차이") and not report["오류"] else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
